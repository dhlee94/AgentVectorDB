"""Excel 파서: 단순 표는 Markdown 표로, 병합/다단 헤더 표는 행 단위 서술형 텍스트로.

[설계 의도]
- 단순 표(헤더 1행, 병합 없음)는 `pandas.to_markdown()`이 가장 정확하고 토큰 효율도 좋습니다.
- 그러나 실무 Excel은 대부분 "보기 좋게" 만든 문서라 병합 셀과 2~3단 헤더가 흔합니다.
  이를 그대로 Markdown 표로 만들면
    (1) 병합 셀 값이 첫 칸에만 남아 나머지 행은 "구분" 값이 빈칸이 되고,
    (2) 청킹으로 표가 잘리면 어떤 열이 어떤 상위 헤더 소속인지 알 수 없게 됩니다.
  그래서 복잡한 표는 병합 셀을 먼저 펼쳐 채운 뒤(forward-fill), 각 행을
  `[행12] 구분: 사무용품, 품목_규격: A4, 가격_단가: 100`처럼 "그 한 줄만 읽어도 완결된"
  서술형 문장으로 변환합니다. 어느 행이 어느 청크로 가든 문맥이 보존됩니다.
- 행 번호는 Excel 실제 행 번호를 사용합니다. 사용자가 인용을 보고 원본 파일에서
  바로 해당 행을 찾을 수 있어야 하기 때문입니다.
"""

from __future__ import annotations

import datetime as dt
import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from rag.exceptions import EncryptedDocumentError, ParserError
from rag.markdown_utils import escape_table_cell
from rag.parsers.base import BaseParser, ParseOutput

logger = logging.getLogger(__name__)

# 암호화된 .xlsx는 ZIP이 아니라 OLE(CFB) 컨테이너로 저장됩니다. 확장자는 xlsx인데
# 시그니처가 OLE면 "손상 파일"이 아니라 "암호 파일"이라고 정확히 알려줄 수 있습니다.
_OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")

# (min_row, min_col, max_row, max_col) — 0-based, 양 끝 포함
MergeRange = Tuple[int, int, int, int]


@dataclass
class SheetGrid:
    """시트 하나를 정규화된 문자열 격자로 표현한 것."""

    name: str
    rows: List[List[str]]
    row_numbers: List[int]  # rows[i]의 Excel 실제 행 번호 (1-based)
    merges: List[MergeRange] = field(default_factory=list)  # rows 좌표계 기준


@dataclass
class SheetResult:
    markdown: str
    mode: str  # "markdown_table" | "narrative"


class ExcelParser(BaseParser):
    supported_extensions = frozenset({".xlsx", ".xlsm", ".xls"})

    # ------------------------------------------------------------------ main
    def _parse(self, file_path: Path) -> ParseOutput:
        warnings: List[str] = []
        if file_path.suffix.lower() == ".xls":
            grids = self._load_xls(file_path, warnings)
        else:
            grids = self._load_xlsx(file_path, warnings)

        sections: List[str] = []
        sheet_modes: Dict[str, str] = {}
        for grid in grids:
            try:
                result = self._sheet_to_markdown(grid)
            except Exception as exc:  # noqa: BLE001
                # 시트 하나의 이상 데이터 때문에 워크북 전체를 버리지 않습니다 (시트 단위 격리).
                warnings.append(f"시트 '{grid.name}' 변환 실패로 건너뜀: {exc}")
                continue
            if result is None:
                warnings.append(f"시트 '{grid.name}'가 비어 있어 건너뜀")
                continue
            # 시트 = "## 시트: 이름" 섹션. 헤더 기반 청킹에서 시트가 자연스러운 Parent 경계가 됩니다.
            sections.append(f"## 시트: {grid.name}\n\n{result.markdown}")
            sheet_modes[grid.name] = result.mode

        # 모든 시트가 비었으면 빈 문자열을 반환 → BaseParser가 EmptyDocumentError로 처리
        metadata = {"sheets": list(sheet_modes), "sheet_modes": sheet_modes}
        return ParseOutput("\n\n".join(sections), "openpyxl+pandas", metadata, warnings)

    # --------------------------------------------------------------- loading
    def _load_xlsx(self, file_path: Path, warnings: List[str]) -> List[SheetGrid]:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException

        with file_path.open("rb") as f:
            if f.read(8) == _OLE_SIGNATURE:
                raise EncryptedDocumentError("암호가 설정된 Excel 파일입니다", source=file_path)

        try:
            # read_only=True가 메모리에는 유리하지만 merged_cells 정보를 제공하지 않습니다.
            # 병합 셀 처리가 이 파서의 존재 이유이므로 일반 모드로 엽니다.
            # data_only=True: 수식 대신 "마지막으로 Excel에서 계산된 값"을 읽습니다.
            #   (주의) 프로그램으로 생성되어 Excel에서 한 번도 열리지 않은 파일은 캐시 값이 없어
            #   수식 셀이 빈칸으로 나옵니다. 이 경우 warnings로 남기지 않고 빈 셀로 처리됩니다.
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
            raise ParserError(f"Excel 파일을 열 수 없습니다(손상 가능성): {exc}", source=file_path) from exc

        grids: List[SheetGrid] = []
        try:
            for ws in wb.worksheets:
                if not hasattr(ws, "merged_cells"):  # 차트 시트 등 데이터가 없는 시트
                    continue
                if ws.sheet_state != "visible" and not self.config.excel_include_hidden_sheets:
                    warnings.append(f"숨김 시트 '{ws.title}' 건너뜀")
                    continue
                grids.append(self._worksheet_to_grid(ws))
        finally:
            wb.close()
        return grids

    def _worksheet_to_grid(self, ws: Any) -> SheetGrid:
        max_row, max_col = ws.max_row or 0, ws.max_column or 0
        rows: List[List[str]] = []
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            rows.append([self._cell_to_str(cell.value, cell.number_format) for cell in row])

        merges: List[MergeRange] = [
            (r.min_row - 1, r.min_col - 1, r.max_row - 1, r.max_col - 1)
            for r in ws.merged_cells.ranges
        ]
        # 병합 셀 펼치기: openpyxl은 병합 영역의 좌상단 셀에만 값을 두고 나머지는 None입니다.
        # 모든 칸에 같은 값을 채워야 "행 단위로 떼어 읽어도" 각 행이 자기 구분값을 갖게 됩니다.
        for r1, c1, r2, c2 in merges:
            if r1 >= len(rows) or c1 >= max_col:
                continue
            value = rows[r1][c1]
            for r in range(r1, min(r2, len(rows) - 1) + 1):
                for c in range(c1, min(c2, max_col - 1) + 1):
                    rows[r][c] = value

        return SheetGrid(
            name=ws.title,
            rows=rows,
            row_numbers=list(range(1, len(rows) + 1)),
            merges=merges,
        )

    def _load_xls(self, file_path: Path, warnings: List[str]) -> List[SheetGrid]:
        """구형 .xls는 openpyxl이 지원하지 않아 pandas(xlrd)로 읽습니다.

        xlrd 경로에서는 병합 정보를 쓰지 않으므로 병합 셀이 있는 .xls는 품질이 떨어집니다.
        운영에서는 .xlsx로 일괄 변환하는 것을 권장합니다.
        """
        try:
            frames = pd.read_excel(file_path, sheet_name=None, header=None, engine="xlrd", dtype=object)
        except ImportError as exc:
            raise ParserError(".xls 파일 처리에는 xlrd 패키지가 필요합니다", source=file_path) from exc
        warnings.append(".xls 형식은 병합 셀 정보를 활용하지 못합니다 (.xlsx 변환 권장)")

        grids: List[SheetGrid] = []
        for name, df in frames.items():
            rows = [[self._cell_to_str(v) for v in row] for row in df.itertuples(index=False)]
            grids.append(SheetGrid(name=str(name), rows=rows, row_numbers=list(range(1, len(rows) + 1))))
        return grids

    # ------------------------------------------------------------ conversion
    def _sheet_to_markdown(self, grid: SheetGrid) -> Optional[SheetResult]:
        grid = self._trim_empty(grid)
        if not grid.rows:
            return None

        preface, grid = self._split_title_rows(grid)
        if not grid.rows:
            # 제목 행만 있는 시트 (예: 표지 시트)
            return SheetResult("\n\n".join(preface), "markdown_table")

        header_count = self._detect_header_rows(grid)
        headers = self._flatten_headers(grid.rows[:header_count])
        body = grid.rows[header_count:]
        body_numbers = grid.row_numbers[header_count:]

        body_has_merge = any(r2 >= header_count for (_, _, r2, _) in grid.merges)
        is_complex = self.config.excel_force_narrative or header_count > 1 or body_has_merge

        parts: List[str] = list(preface)
        if not body:
            # 헤더만 있는 표 — 헤더 자체라도 정보로 남깁니다.
            parts.append(", ".join(headers))
            return SheetResult("\n\n".join(parts), "markdown_table")

        if is_complex:
            parts.append(self._to_narrative(headers, body, body_numbers))
            return SheetResult("\n\n".join(parts), "narrative")

        parts.append(self._to_markdown_table(headers, body))
        return SheetResult("\n\n".join(parts), "markdown_table")

    @staticmethod
    def _trim_empty(grid: SheetGrid) -> SheetGrid:
        """완전히 빈 행/열을 제거하고 병합 좌표를 새 좌표계로 옮깁니다.

        Excel은 서식만 지정된 빈 칸도 사용 영역(max_row/max_column)에 포함시키므로,
        제거하지 않으면 빈 열 수십 개짜리 표나 빈 행 수천 개가 생깁니다.
        """
        keep_rows = [i for i, row in enumerate(grid.rows) if any(v for v in row)]
        if not keep_rows:
            return SheetGrid(grid.name, [], [], [])
        width = max(len(r) for r in grid.rows)
        keep_cols = [c for c in range(width) if any(c < len(grid.rows[i]) and grid.rows[i][c] for i in keep_rows)]

        row_map = {old: new for new, old in enumerate(keep_rows)}
        col_map = {old: new for new, old in enumerate(keep_cols)}
        rows = [[grid.rows[i][c] if c < len(grid.rows[i]) else "" for c in keep_cols] for i in keep_rows]
        numbers = [grid.row_numbers[i] for i in keep_rows]

        merges: List[MergeRange] = []
        for r1, c1, r2, c2 in grid.merges:
            rs = [row_map[r] for r in range(r1, r2 + 1) if r in row_map]
            cs = [col_map[c] for c in range(c1, c2 + 1) if c in col_map]
            if rs and cs:
                merges.append((min(rs), min(cs), max(rs), max(cs)))
        return SheetGrid(grid.name, rows, numbers, merges)

    @staticmethod
    def _split_title_rows(grid: SheetGrid) -> Tuple[List[str], SheetGrid]:
        """표 위에 있는 제목/단위 행("2025년 견적서", "(단위: 원)")을 분리합니다.

        이런 행을 헤더로 오인하면 모든 열 이름이 "2025년 견적서"가 되어 버립니다.
        값이 한 종류뿐인 행(병합된 제목 포함)을 표 앞의 문단으로 빼내면, 청커가 이를
        "표 캡션"으로 인식해 잘린 표 조각마다 붙여 줍니다.
        """
        width = len(grid.rows[0]) if grid.rows else 0
        if width <= 1:
            return [], grid  # 한 열짜리 시트는 모든 행이 "값 한 종류"라 판별 불가

        preface: List[str] = []
        idx = 0
        # 최대 3행까지만 제목으로 인정 (데이터 행을 제목으로 먹어버리는 오판 방지)
        while idx < min(3, len(grid.rows) - 1):
            distinct = {v for v in grid.rows[idx] if v}
            if len(distinct) != 1:
                break
            preface.append(distinct.pop())
            idx += 1

        if idx == 0:
            return [], grid
        merges = [
            (r1 - idx, c1, r2 - idx, c2) for (r1, c1, r2, c2) in grid.merges if r2 >= idx
        ]
        merges = [(max(r1, 0), c1, r2, c2) for (r1, c1, r2, c2) in merges]
        return preface, SheetGrid(grid.name, grid.rows[idx:], grid.row_numbers[idx:], merges)

    def _detect_header_rows(self, grid: SheetGrid) -> int:
        """다단 헤더 깊이를 추정합니다.

        규칙: 헤더 행 안에 "가로로 병합된 셀"이 있으면 그것은 상위 그룹 헤더
        (예: '가격' 아래 '단가'·'합계')이므로 다음 행도 헤더에 포함합니다.
        """
        header_rows = 1
        limit = min(self.config.excel_max_header_rows, len(grid.rows) - 1)
        while header_rows < limit:
            last = header_rows - 1
            next_row = grid.rows[header_rows]
            has_group_header = False
            for r1, c1, r2, c2 in grid.merges:
                if not (r1 == r2 == last and c2 > c1):
                    continue
                # 가로 병합만으로는 부족합니다. 단순히 넓게 병합한 '비고' 칸일 수도 있기 때문입니다.
                # 바로 아래 행의 해당 구간에 서로 다른 하위 헤더가 있어야 그룹 헤더로 인정합니다.
                below = [next_row[c] for c in range(c1, c2 + 1) if c < len(next_row)]
                if all(below) and len(set(below)) > 1:
                    has_group_header = True
                    break
            if not has_group_header:
                break
            header_rows += 1
        return header_rows

    @staticmethod
    def _flatten_headers(header_rows: List[List[str]]) -> List[str]:
        """다단 헤더를 '상위_하위' 한 줄로 평탄화하고, 빈/중복 이름을 보정합니다."""
        width = max(len(r) for r in header_rows)
        names: List[str] = []
        seen: Dict[str, int] = {}
        for c in range(width):
            parts: List[str] = []
            for row in header_rows:
                value = row[c] if c < len(row) else ""
                # 세로 병합(예: '품목'이 2행에 걸침)은 forward-fill로 같은 값이 반복되므로 연속 중복 제거
                if value and (not parts or parts[-1] != value):
                    parts.append(value)
            name = "_".join(parts) if parts else f"열{c + 1}"
            # 같은 이름의 열이 두 개면 "가격: 100, 가격: 200"처럼 모호해지므로 번호를 붙입니다.
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 1
            names.append(name)
        return names

    @staticmethod
    def _to_markdown_table(headers: List[str], body: List[List[str]]) -> str:
        df = pd.DataFrame(
            [[escape_table_cell(v) for v in row] for row in body],
            columns=[escape_table_cell(h) for h in headers],
        )
        # disable_numparse=True가 없으면 tabulate가 "00123"(제품코드) → 123,
        # "1.50"(규격) → 1.5로 숫자 변환해 버립니다. 원문 그대로를 보존해야 합니다.
        table = df.to_markdown(index=False, disable_numparse=True)
        # tabulate는 열 정렬을 위해 공백을 채워 넣는데, 긴 셀 하나 때문에 모든 행이
        # 수십 칸씩 부풀어 청크 예산과 토큰을 낭비합니다. 연속 공백을 하나로 줄입니다.
        return "\n".join(re.sub(r" {2,}", " ", line) for line in table.split("\n"))

    @staticmethod
    def _to_narrative(headers: List[str], body: List[List[str]], row_numbers: List[int]) -> str:
        """각 행을 '[행N] 열이름: 값, ...' 한 줄로 변환합니다 (빈 값은 생략)."""
        lines: List[str] = []
        for row, number in zip(body, row_numbers):
            pairs = [f"{h}: {v}" for h, v in zip(headers, row) if v]
            if pairs:
                lines.append(f"[행{number}] " + ", ".join(pairs))
        # 행마다 빈 줄로 구분 → 청커가 각 행을 독립 문단(블록)으로 다뤄 행 중간에서 자르지 않습니다.
        return "\n\n".join(lines)

    # ---------------------------------------------------------------- values
    @staticmethod
    def _cell_to_str(value: Any, number_format: str = "General") -> str:
        """셀 값을 사람이 Excel에서 보는 모양에 가깝게 문자열화합니다."""
        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, dt.datetime):
            # 날짜만 있는 셀도 openpyxl은 datetime(…, 0, 0)으로 줍니다. 00:00:00은 잡음이므로 제거.
            if value.time() == dt.time(0, 0):
                return value.date().isoformat()
            return value.strftime("%Y-%m-%d %H:%M")
        if isinstance(value, dt.date):
            return value.isoformat()
        if isinstance(value, dt.time):
            return value.strftime("%H:%M")
        if isinstance(value, (int, float)):
            if isinstance(value, float) and value != value:  # NaN (pandas .xls 경로)
                return ""
            # 퍼센트 서식 셀은 0.15로 저장되어 있습니다. 그대로 두면 LLM이 "할인율 0.15"를
            # 15%로 해석하지 못해 오답을 냅니다.
            if isinstance(number_format, str) and "%" in number_format:
                pct = value * 100
                return f"{int(pct)}%" if float(pct).is_integer() else f"{round(pct, 4)}%"
            if isinstance(value, float):
                # 부동소수 잡음(0.30000000000000004) 제거, 정수형 실수(100.0)는 정수로
                return str(int(value)) if value.is_integer() else str(round(value, 10))
            return str(value)
        text = str(value)
        # 셀 내 줄바꿈(Alt+Enter)은 행 단위 서술형/Markdown 표를 깨뜨리므로 공백으로 치환
        return re.sub(r"\s+", " ", text).strip()
