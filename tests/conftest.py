"""공용 픽스처.

원칙
- 실제 Claude API와 bge-m3 모델을 쓰지 않습니다 (비용·네트워크·2GB 다운로드 없이 수 초 안에 실행).
  임베딩은 DeterministicFakeEmbedding, LLM은 호출 인자를 기록하는 가짜 클라이언트로 대체합니다.
- 샘플 문서(PDF/Excel/Word)는 바이너리를 커밋하지 않고 테스트 시작 시 코드로 생성합니다.
- 각 테스트는 임시 디렉터리에서 실행해 저장소의 config.yaml / .env(실제 API 키)를 읽지 않습니다.
"""

from __future__ import annotations

import datetime as dt
import types
from pathlib import Path
from typing import Any, Dict, List, Optional

import pytest


# ---------------------------------------------------------------------------
# 샘플 문서 생성
# ---------------------------------------------------------------------------
def _make_excel(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "단가표"  # 단순 표 → Markdown 표
    ws.append(["품목코드", "품목", "단가"])
    for i in range(1, 6):
        ws.append([f"00{i}", f"품목{i}", 100 * i])

    ws2 = wb.create_sheet("견적")  # 제목 행 + 2단 헤더 + 본문 병합 → 서술형
    ws2["A1"] = "2025년 견적서"
    ws2.merge_cells("A1:E1")
    ws2["A2"] = "구분"
    ws2.merge_cells("A2:A3")
    ws2["B2"] = "품목"
    ws2.merge_cells("B2:B3")
    ws2["C2"] = "가격"
    ws2.merge_cells("C2:D2")
    ws2["C3"] = "단가"
    ws2["D3"] = "합계"
    ws2["E2"] = "할인율"
    ws2.merge_cells("E2:E3")
    rows = [("사무용품", "A4용지", 100, 300, 0.15), (None, "볼펜", 50, 500, 0.1), ("가구", "의자", 90000, 90000, 0)]
    for r, (a, b, c, d, e) in enumerate(rows, start=4):
        ws2.cell(r, 1, a)
        ws2.cell(r, 2, b)
        ws2.cell(r, 3, c)
        ws2.cell(r, 4, d)
        ws2.cell(r, 5, e).number_format = "0%"
    ws2.merge_cells("A4:A5")
    ws2["B7"] = dt.datetime(2025, 7, 15)

    wb.create_sheet("빈시트")
    hidden = wb.create_sheet("숨김")
    hidden["A1"] = "x"
    hidden.sheet_state = "hidden"
    wb.save(path)


def _make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("1장 개요", 1)
    d.add_paragraph("이 문서는 테스트용 사내 규정입니다.")
    d.add_heading("1.1 적용 범위", 2)
    d.add_paragraph("모든 임직원에게 적용된다.")
    d.add_heading("2장 단가", 1)
    d.add_paragraph("표 1. 부품 단가표")
    t = d.add_table(rows=1, cols=3)
    for cell, text in zip(t.rows[0].cells, ["항목", "가격", "비고"]):
        cell.text = text
    for i in range(40):
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = f"부품-{i:03d}", str(1000 + i), "재고 있음 | 긴급"
    d.add_paragraph("표 아래 설명 문단.")
    d.save(path)


def _make_pdf(path: Path) -> None:
    import fitz

    pdf = fitz.open()
    for p in range(3):
        page = pdf.new_page()
        page.insert_text((72, 72), f"Chapter {p + 1} Pricing", fontsize=20)
        for i in range(20):
            page.insert_text((72, 112 + 16 * i), f"Line {i} of page {p + 1}: unit price is {100 + i} won.", fontsize=10)
    pdf.save(path)


@pytest.fixture(scope="session")
def samples(tmp_path_factory) -> Dict[str, Path]:
    import fitz

    root = tmp_path_factory.mktemp("samples")
    files = {
        "xlsx": root / "견적_테스트.xlsx",
        "docx": root / "규정_테스트.docx",
        "pdf": root / "manual.pdf",
        "locked": root / "locked.pdf",
        "blank": root / "blank.pdf",
        "empty": root / "empty.pdf",
        "lockfile": root / "~$견적_테스트.xlsx",
        "txt": root / "notes.txt",
    }
    _make_excel(files["xlsx"])
    _make_docx(files["docx"])
    _make_pdf(files["pdf"])
    enc = fitz.open()
    enc.new_page().insert_text((72, 72), "secret")
    enc.save(files["locked"], encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    blank = fitz.open()
    blank.new_page()
    blank.save(files["blank"])
    files["empty"].write_bytes(b"")
    files["lockfile"].write_bytes(b"x")
    files["txt"].write_text("hi", encoding="utf-8")
    return files


# ---------------------------------------------------------------------------
# 가짜 LLM / 파이프라인
# ---------------------------------------------------------------------------
class FakeMessages:
    """client.beta.messages.create(**kwargs) 대역. 마지막 호출 인자와 호출 횟수를 기록합니다."""

    def __init__(self) -> None:
        self.reply = "답변입니다 [1]."
        self.stop_reason = "end_turn"
        self.raise_exc: Optional[Exception] = None
        self.calls: List[Dict[str, Any]] = []

    @property
    def last(self) -> Dict[str, Any]:
        return self.calls[-1]

    def create(self, **kwargs: Any) -> Any:
        self.calls.append(kwargs)
        if self.raise_exc is not None:
            raise self.raise_exc
        usage = types.SimpleNamespace(input_tokens=100, output_tokens=20, cache_read_input_tokens=0, iterations=None)
        return types.SimpleNamespace(
            content=[
                types.SimpleNamespace(type="thinking", thinking=""),
                types.SimpleNamespace(type="text", text=self.reply),
            ],
            stop_reason=self.stop_reason,
            stop_details=None,
            model=kwargs["model"],
            usage=usage,
            _request_id="req_test",
        )


@pytest.fixture
def fake_llm() -> FakeMessages:
    return FakeMessages()


@pytest.fixture(autouse=True)
def isolate_from_repo(tmp_path, monkeypatch):
    """저장소의 config.yaml / .env(실제 API 키)가 테스트에 섞이지 않도록 임시 디렉터리에서 실행."""
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("RAG_CONFIG", raising=False)


def make_pipeline(store_dir: Path, fake_llm: FakeMessages, embedding_model: str = "fake-256"):
    from langchain_core.embeddings import DeterministicFakeEmbedding

    from rag.config import PipelineConfig
    from rag.pipeline import RAGPipeline

    cfg = PipelineConfig()
    cfg.store.persist_dir = str(store_dir)
    cfg.store.embedding_model = embedding_model
    client = types.SimpleNamespace(beta=types.SimpleNamespace(messages=fake_llm))
    return RAGPipeline(cfg, embeddings=DeterministicFakeEmbedding(size=256), llm_client=client)


@pytest.fixture
def inbox(tmp_path, samples) -> Path:
    """xlsx / docx / pdf / 암호 PDF가 들어 있는 인제스트용 폴더."""
    import shutil

    folder = tmp_path / "inbox"
    folder.mkdir()
    for key in ("xlsx", "docx", "pdf", "locked"):
        shutil.copy(samples[key], folder / samples[key].name)
    return folder


@pytest.fixture
def indexed(tmp_path, inbox, fake_llm):
    """inbox를 인제스트한 파이프라인."""
    pipe = make_pipeline(tmp_path / "store", fake_llm)
    report = pipe.ingest(str(inbox))
    assert len(report.indexed) == 3, report.summary()
    return pipe
